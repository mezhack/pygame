import docx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Inches


def set_cell_background(cell, hex_color):
    """Define a cor de fundo de uma célula."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define as margens internas (padding) de uma célula."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [
        ('top', top),
        ('bottom', bottom),
        ('left', left),
        ('right', right),
    ]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_teacher_guide_doc(filename='Guia_Professor_Gabarito.docx'):
    doc = Document()

    # Configuração de Margens da Página
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos Base
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # ---------------------------------------------------------
    # CABEÇALHO DO DOCUMENTO
    # ---------------------------------------------------------
    title = doc.add_paragraph()
    title_run = title.add_run('Plano de Aula & Gabarito Orientado')
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Azul Escuro
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run(
        'Material de Apoio Pedagógico ao Docente | Instruções e Respostas Comentadas'
    )
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    subtitle.paragraph_format.space_after = Pt(18)

    # Linha Divisória
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(18)
    p_line_run = p_line.add_run(
        '_________________________________________________________________________________'
    )
    p_line_run.font.color.rgb = RGBColor(0xD3, 0xD3, 0xD3)

    # ---------------------------------------------------------
    # SEÇÃO 1: ORIENTAÇÕES PEDAGÓGICAS
    # ---------------------------------------------------------
    h1 = doc.add_paragraph()
    h1_run = h1.add_run('1. Orientações Pedagógicas para a Condução da Aula')
    h1_run.font.name = 'Arial'
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    # Caixa Destaque de Nota Pedagógica (Tabela de 1 célula)
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout.autofit = False
    callout.columns[0].width = Inches(6.5)

    c_cell = callout.cell(0, 0)
    set_cell_background(c_cell, 'F0F4F8')  # Fundo azul claro
    set_cell_margins(c_cell, top=140, bottom=140, left=200, right=200)

    # Borda esquerda grossa
    tcPr = c_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    c_bold = cp.add_run('💡 Objetivo da Abordagem:')
    c_bold.font.bold = True
    c_bold.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    cp2 = c_cell.add_paragraph()
    cp2.paragraph_format.space_after = Pt(0)
    cp2_text = cp2.add_run(
        'Incentivar a autonomia, o raciocínio crítico e a resolução ativa de problemas pelos estudantes. '
        'Evite fornecer as respostas prontas imediatamente; utilize o método de mediação instrucional.'
    )
    cp2_text.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Passos Sequenciais
    steps = [
        (
            'Momento Inicial (Sensibilização & Contextualização):',
            ' Inicie a aula fazendo perguntas provocativas sobre o tema central para levantar os conhecimentos prévios da turma. Apresente o desafio do dia.',
        ),
        (
            'Desenvolvimento (Prática Guiada & Trabalho Individual/Grupo):',
            ' Permita que os alunos tentem resolver as atividades autonomamente antes de intervir. Circule pela sala observando dificuldades recorrentes.',
        ),
        (
            'Intervenção Mediada (Suporte Pontual):',
            ' Quando surgirem dúvidas, faça perguntas norteadoras em vez de dar a resposta final. Ajude o aluno a identificar a lógica do erro.',
        ),
        (
            'Encerramento (Fechamento Cognitivo):',
            ' Reserve os últimos 15 minutos para fazer a correção coletiva com a turma no quadro ou projetor, alinhando as respostas ao gabarito oficial.',
        ),
    ]

    for title_step, desc in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        run_t = p.add_run(title_step)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        run_d = p.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # SEÇÃO 2: GABARITO E RESPOSTAS COMENTADAS
    # ---------------------------------------------------------
    h2 = doc.add_paragraph()
    h2_run = h2.add_run('2. Respostas Comentadas e Critérios de Avaliação')
    h2_run.font.name = 'Arial'
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(8)

    # Tabela de Gabarito
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Definir larguras das colunas
    col_widths = [Inches(1.0), Inches(3.0), Inches(2.5)]

    # Cabeçalho da Tabela
    hdr_cells = table.rows[0].cells
    headers = ['Questão', 'Resposta Esperada / Gabarito', 'Critérios de Correção & Dicas']
    for i, title_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], '1B365D')
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title_text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)

    # Dados do Gabarito (Exemplo estruturado - Substituir ou expandir conforme necessário)
    dados_gabarito = [
        (
            'Questão 01',
            'Resposta conceitual correta detalhando o funcionamento principal da estrutura estudada.',
            'Considerar correto se o aluno citar ao menos 2 conceitos-chave. Aceitar variações sinonímicas.',
        ),
        (
            'Questão 02',
            'A aplicação prática deve demonstrar a sequência lógica de passos (Passo A → Passo B → Passo C).',
            'Verificar a consistência do raciocínio e a correta aplicação da fórmula/regra.',
        ),
        (
            'Questão 03',
            'Análise crítica ou justificativa argumentativa esperada.',
            'Valorizar a clareza na articulação das ideias e a fundamentação teórica apresentada.',
        ),
    ]

    for row_idx, (q, r, c) in enumerate(dados_gabarito):
        row_cells = table.add_row().cells
        bg_color = 'F9FAFB' if row_idx % 2 == 0 else 'FFFFFF'

        for i, text in enumerate([q, r, c]):
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(
                row_cells[i], top=100, bottom=100, left=120, right=120
            )
            p = row_cells[i].paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)

            if i == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # ---------------------------------------------------------
    # SEÇÃO 3: CRONOGRAMA DA AULA
    # ---------------------------------------------------------
    h3 = doc.add_paragraph()
    h3_run = h3.add_run('3. Sugestão de Cronograma de Tempo (50 min)')
    h3_run.font.name = 'Arial'
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(8)

    schedule_table = doc.add_table(rows=1, cols=2)
    schedule_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    schedule_table.autofit = False

    s_widths = [Inches(1.5), Inches(5.0)]
    s_hdr = schedule_table.rows[0].cells
    s_hdr[0].width, s_hdr[1].width = s_widths[0], s_widths[1]

    set_cell_background(s_hdr[0], '4A5568')
    set_cell_background(s_hdr[1], '4A5568')

    for cell, text in zip(s_hdr, ['Tempo Estimado', 'Atividade']):
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)

    cronograma = [
        ('10 minutos', 'Introdução do tema, motivação inicial e alinhamento dos objetivos.'),
        ('25 minutos', 'Resolução individual ou em duplas das atividades pelos estudantes.'),
        ('15 minutos', 'Debate coletivo, tirada de dúvidas e fechamento com o gabarito.'),
    ]

    for row_idx, (tempo, atividade) in enumerate(cronograma):
        r_cells = schedule_table.add_row().cells
        bg_color = 'F9FAFB' if row_idx % 2 == 0 else 'FFFFFF'

        for i, text in enumerate([tempo, atividade]):
            r_cells[i].width = s_widths[i]
            set_cell_background(r_cells[i], bg_color)
            set_cell_margins(
                r_cells[i], top=90, bottom=90, left=120, right=120
            )
            p = r_cells[i].paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if i == 0:
                run.font.bold = True

    # Salvar Arquivo
    doc.save(filename)
    print(f'Documento gerado com sucesso: {filename}')


if __name__ == '__main__':
    create_teacher_guide_doc()